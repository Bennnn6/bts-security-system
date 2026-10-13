import os
import sys
import subprocess

# 1. Otomatis install python-docx jika belum terpasang
try:
    import docx
except ImportError:
    print("[INFO] Mengunduh & memasang library python-docx...")
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])
    import docx

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def set_cell_background(cell, fill_hex):
    """Memberikan warna latar belakang pada header tabel"""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def generate_docs():
    print("[INFO] Sedang menyusun isi dokumen Word...")
    doc = Document()
    
    # Margin Halaman (1 Inci / 2.54 cm)
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # ==========================================
    # COVER HALAMAN JUDUL
    # ==========================================
    tp = doc.add_paragraph()
    tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tp.paragraph_format.space_before = Pt(80)
    tp.paragraph_format.space_after = Pt(12)
    
    r_title = tp.add_run("DOKUMENTASI SISTEM & SPESIFIKASI LENGKAP\nBTS SMART ACCESS SECURITY SYSTEM")
    r_title.font.name = 'Arial'
    r_title.font.size = Pt(22)
    r_title.font.bold = True
    r_title.font.color.rgb = RGBColor(16, 44, 87) # Dark Navy

    sp = doc.add_paragraph()
    sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sp.paragraph_format.space_after = Pt(140)
    r_sub = sp.add_run("Sistem Pengendalian dan Pemantauan Akses Fisik Shelter BTS Menara Telekomunikasi\nBerbasis 3-Factor Authentication (Work Order, Face Recognition, & TOTP)")
    r_sub.font.name = 'Arial'
    r_sub.font.size = Pt(12)
    r_sub.font.italic = True
    r_sub.font.color.rgb = RGBColor(90, 90, 90)

    mp = doc.add_paragraph()
    mp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_meta = mp.add_run("Disusun untuk:\nDokumentasi Teknis Proyek Tugas Akhir / Laporan Sistem Keamanan BTS\nVersi 1.0 — Tahun 2026")
    r_meta.font.name = 'Arial'
    r_meta.font.size = Pt(11)
    r_meta.font.color.rgb = RGBColor(60, 60, 60)

    doc.add_page_break()

    # ==========================================
    # BAB 1: PENDAHULUAN
    # ==========================================
    doc.add_heading('BAB 1: PENDAHULUAN & LATAR BELAKANG SISTEM', level=1)
    doc.add_paragraph(
        "Base Transceiver Station (BTS) merupakan infrastruktur telekomunikasi vital yang menjamin kelangsungan konektivitas "
        "jaringan seluler bagi masyarakat luas dan sektor industri. Di dalam shelter dan area pagar BTS tersimpan aset "
        "bernilai investasi sangat tinggi, meliputi modul Baseband Unit (BBU), Remote Radio Unit (RRU), rectifier catu daya, "
        "modul transmisi optik/microwave, baterai cadangan (Lithium/Lead-Acid) berkapasitas besar, dan generator set (Genset)."
    )
    doc.add_heading('1.1 Masalah pada Pengamanan Konvensional', level=2)
    doc.add_paragraph("• Kunci Fisik / Gembok Manual: Sangat rawan digandakan, hilang, atau dirusak paksa dengan cairan kimia dan gergaji.", style='List Bullet')
    doc.add_paragraph("• Akses Tanpa Otorisasi: Teknisi atau pihak luar dapat memasuki area menara tanpa jadwal kerja resmi di luar jam operasional.", style='List Bullet')
    doc.add_paragraph("• Ketiadaan Jejak Audit (Audit Trail): Tim NOC tidak memiliki catatan otomatis mengenai siapa yang membuka pintu dan kapan pintu ditutup.", style='List Bullet')
    doc.add_paragraph("• Keterlambatan Respon Alarm: Kasus pencurian baterai atau sabotase kabel transmisi sering kali terlambat disadari.", style='List Bullet')

    doc.add_heading('1.2 Solusi yang Diterapkan', level=2)
    doc.add_paragraph(
        "Aplikasi web ini menghadirkan ekosistem keamanan cerdas terpadu yang menggantikan kunci mekanis dengan kontrol digital "
        "bertingkat (3-Factor Authentication), pencatatan audit log otomatis, kontrol hardware relay, serta sistem peringatan dini (early warning alarm)."
    )

    # ==========================================
    # BAB 2: ARSITEKTUR TEKNOLOGI
    # ==========================================
    doc.add_heading('BAB 2: ARSITEKTUR TEKNOLOGI & SPESIFIKASI SISTEM', level=1)
    doc.add_paragraph(
        "Sistem ini dibangun dengan arsitektur modular berbasis Flask (Python) yang mengintegrasikan komponen dashboard web, "
        "kecerdasan buatan (computer vision), modul kriptografi waktu nyata, dan pengendali perangkat keras:"
    )

    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    hdr[0].text, hdr[1].text, hdr[2].text = "Komponen", "Teknologi", "Fungsi & Peran Utama"
    for c in hdr:
        set_cell_background(c, "102C57")
        c.paragraphs[0].runs[0].font.bold = True
        c.paragraphs[0].runs[0].font.color.rgb = RGBColor(255, 255, 255)

    data_tech = [
        ("Backend Framework", "Python / Flask", "Pusat logika bisnis, session management, routing halaman web, dan RESTful API."),
        ("Basis Data & ORM", "SQLite / SQLAlchemy / Flask-Migrate", "Penyimpanan data relasional, pencatatan log transaksi, dan migrasi skema database."),
        ("Autentikasi User", "Flask-Login & Werkzeug", "Manajemen sesi login web, enkripsi hashing password, dan Role-Based Access Control (RBAC)."),
        ("Biometrik Wajah", "OpenCV / Face Recognition (128D)", "Mendeteksi wajah teknisi di kamera, ekstraksi facial landmark, dan perbandingan vektor biometrik."),
        ("Dynamic OTP", "PyOTP (RFC 6238 TOTP)", "Memvalidasi kode 6 digit dinamis yang berganti setiap 30 detik untuk mencegah serangan anti-replay."),
        ("Kontrol Fisik (Relay)", "GPIO / Serial / HTTP Webhook", "Memicu arus listrik untuk membuka kunci solenoid/magnetic lock pintu shelter BTS."),
        ("Sistem Notifikasi", "Telegram Bot / Webhook Gateway", "Mengirimkan peringatan instan ke tim NOC saat terjadi pembobolan atau insiden pintu."),
        ("Antarmuka Frontend", "HTML5, CSS3, JavaScript, Jinja2, PWA", "Dashboard monitoring NOC dan portal ramah mobile bagi teknisi di lapangan.")
    ]

    for k, t, f in data_tech:
        row = table.add_row()
        row.cells[0].text, row.cells[1].text, row.cells[2].text = k, t, f
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # ==========================================
    # BAB 3: MEKANISME 3-FACTOR AUTHENTICATION
    # ==========================================
    doc.add_heading('BAB 3: MEKANISME KEAMANAN 3-FACTOR AUTHENTICATION (3FA)', level=1)
    doc.add_paragraph(
        "Untuk menjamin pintu shelter hanya terbuka bagi personel yang sah pada waktu dan lokasi yang tepat, sistem menerapkan sistem interlocking 3 lapis:"
    )
    doc.add_heading('Faktor 1: Work Order / Surat Tugas Sah (Authorization Layer)', level=2)
    doc.add_paragraph(
        "Teknisi wajib memiliki tiket kerja berstatus APPROVED. Sistem memverifikasi kecocokan site_id dengan lokasi fisik teknisi, "
        "serta memeriksa apakah waktu kunjungan berada dalam rentang valid (start_time s/d end_time)."
    )
    doc.add_heading('Faktor 2: Face Recognition Biometrik (Identity Layer)', level=2)
    doc.add_paragraph(
        "Kamera mengambil foto wajah teknisi secara langsung di pintu gerbang/shelter. Algoritma membandingkan vektor wajah dengan database "
        "profil resmi teknisi. Akses hanya lolos jika tingkat kemiripan (similarity score) memenuhi ambang batas yang ditetapkan (>= 85%)."
    )
    doc.add_heading('Faktor 3: Time-based One-Time Password (Token Layer)', level=2)
    doc.add_paragraph(
        "Teknisi menginputkan 6-digit kode OTP dinamis dari aplikasi smartphone mereka. Kode ini kedaluwarsa setiap 30 detik, "
        "mencegah serangan pemalsuan token maupun penyadapan (anti-replay)."
    )

    # ==========================================
    # BAB 4: STRUKTUR FILE & MODUL
    # ==========================================
    doc.add_heading('BAB 4: STRUKTUR DIREKTORI & PENJELASAN MODUL KODE', level=1)
    files = [
        ("run.py", "Entry point aplikasi untuk menjalankan server web Flask lokal pada port default 5000."),
        ("config.py", "Pusat konfigurasi variabel lingkungan: secret key, database URI, delay relay gerbang, dan toleransi face match."),
        ("app/__init__.py", "Inisialisasi aplikasi Flask, konfigurasi database SQLAlchemy, Flask-Login, migrasi, dan registrasi blueprint rute."),
        ("app/models.py", "Mendefinisikan entitas basis data: User (teknisi/operator), Site (menara BTS), WorkOrder (surat tugas), AccessLog (audit log), dan AlarmEvent."),
        ("app/access_control.py", "Otak logika keamanan yang mengorkestrasi verifikasi 3FA dan mengirim instruksi buka kunci ke relay hardware."),
        ("app/face_recognition_module.py", "Menangani proses ekstraksi landmark wajah dan penghitungan jarak kemiripan (Euclidean distance/cosine similarity)."),
        ("app/totp_module.py", "Menghasilkan dan memverifikasi secret key serta kode OTP berbasis waktu."),
        ("app/notifications.py", "Mengirimkan alert otomatis ke Telegram/NOC ketika terjadi pelanggaran akses atau pintu dibiarkan terbuka."),
        ("test_system.py", "Skrip pengujian otomatis komprehensif untuk memeriksa integritas database, model relasional, dan seluruh endpoint API.")
    ]
    for fn, fd in files:
        p = doc.add_paragraph()
        r = p.add_run(f"• {fn}: ")
        r.bold = True
        r.font.color.rgb = RGBColor(16, 44, 87)
        p.add_run(fd)

    # ==========================================
    # BAB 5: FITUR ANTARMUKA WEB
    # ==========================================
    doc.add_heading('BAB 5: ANTARMUKA WEB (USER INTERFACE) & FITUR OPERASIONAL', level=1)
    doc.add_paragraph("• Dashboard (dashboard.html): Monitoring status site aktif, ringkasan statistik akses hari ini, grafik kunjungan, dan feed aktivitas real-time.")
    doc.add_paragraph("• Access Logs (access_logs.html): Rekaman audit log lengkap yang mencantumkan nama teknisi, waktu presisi, status berhasil/ditolak, alasan penolakan, serta foto tangkapan kamera.")
    doc.add_paragraph("• Site Management (sites.html): Pendataan menara BTS, koordinat geografis (latitude/longitude), dan status operasional perangkat pintu.")
    doc.add_paragraph("• Work Order Management (work_orders.html): Pembuatan, peninjauan, dan persetujuan surat tugas teknisi sebelum kunjungan lapangan.")
    doc.add_paragraph("• User Management (users.html): Pengelolaan akun pengguna, hak akses berbasis peran (RBAC), serta pendaftaran data wajah biometrik.")

    # ==========================================
    # BAB 6: PENANGANAN ANOMALI
    # ==========================================
    doc.add_heading('BAB 6: PENANGANAN ANOMALI & SISTEM PERINGATAN DINI', level=1)
    doc.add_paragraph("• Door Forced Open (Pintu Dibobol Paksa): Jika sensor pintu mendeteksi kontak terlepas tanpa instruksi verifikasi 3FA, sirene lokal aktif dan notifikasi darurat langsung dikirim ke tim keamanan NOC.")
    doc.add_paragraph("• Door Held Open (Pintu Lupa Ditutup): Jika pintu tidak ditutup kembali melebihi batas waktu aman (misal > 3 menit), sistem memicu alarm pengingat ke teknisi dan dispatcher.")
    doc.add_paragraph("• Anti-Spoofing & Blacklist: Kegagalan pencocokan wajah berulang kali otomatis membekukan akses sementara dan menyimpan bukti foto upaya pembobolan.")

    # ==========================================
    # BAB 7: PANDUAN MENJALANKAN SISTEM
    # ==========================================
    doc.add_heading('BAB 7: PANDUAN MENJALANKAN & MENGUJI SISTEM', level=1)
    doc.add_paragraph("1. Aktifkan virtual environment: .\\venv\\Scripts\\activate")
    doc.add_paragraph("2. Pasang dependensi aplikasi: pip install -r requirements.txt")
    doc.add_paragraph("3. Jalankan pengujian otomatis: python test_system.py")
    doc.add_paragraph("4. Jalankan server Flask: python run.py (Akses melalui peramban web pada http://localhost:5000)")

    output_filename = "Dokumentasi_Lengkap_Sistem_Keamanan_BTS.docx"
    doc.save(output_filename)
    
    full_path = os.path.abspath(output_filename)
    print("\n" + "="*60)
    print(f"[BERHASIL] File Word sukses dibuat di:\n{full_path}")
    print("="*60 + "\n")
    
    # Otomatis buka Word di layar Anda
    os.system(f'start "" "{full_path}"')

if __name__ == "__main__":
    generate_docs()